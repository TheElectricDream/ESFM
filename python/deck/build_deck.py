"""Builds the ESFM pipeline deck on the SRL template.

Every number shown comes from the facts_*.json files written by the figure
scripts, i.e. from the pipeline run in the cache. Set ESFM_PREVIEW=1 to stamp
every slide (used only for the synthetic smoke test).
"""
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(os.path.dirname(HERE), 'figures'))
from pptx_helpers import *  # noqa
from pptx.util import Inches, Pt
import eq as EQ

FIG = os.environ.get('ESFM_FIGOUT', os.path.join(os.path.dirname(HERE), 'figures', 'out'))
TEMPLATE = os.environ.get('ESFM_TEMPLATE', os.path.join(HERE, 'SRL_Slide_Template.pptx'))
OUTFILE = os.environ.get('ESFM_DECK', os.path.join(HERE, 'ESFM_pipeline.pptx'))
PREVIEW = os.environ.get('ESFM_PREVIEW') == '1'

fx = {}
for n in ('overview', 'front', 'warm', 'back'):
    p = os.path.join(FIG, f'facts_{n}.json')
    if os.path.exists(p):
        d = json.load(open(p))
        fx.update(d if n == 'overview' else {k: v for k, v in d.items()})
O = fx  # overview facts at top level; per-figure facts under their function names
g = lambda k, sub, default=float('nan'): fx.get(k, {}).get(sub, default)

# ------------------------------------------------------------------ geometry (inches)
FX, FY, FW, FH = 0.5, 1.18, 5.4, 3.35          # figure, left
CX, CW = 6.1, 3.3                              # right column
WX, WW = 0.8, 8.4                              # wide figure
EQ_PT = 16


def fig_paths(name):
    return os.path.join(FIG, name + '.svg'), os.path.join(FIG, name + '.png')


def place_eq(slide, key, latex, x, y, maxw, pt=EQ_PT):
    w, h = EQ.render(key, latex, pt)
    if w > maxw:
        h *= maxw / w; w = maxw
    s, p = os.path.join(EQ.OUT, key + '.svg'), os.path.join(EQ.OUT, key + '.png')
    add_svg(slide, s, p, x, y, w=w, h=h, name='Equation ' + key)
    return h


def stamp(slide):
    if PREVIEW:
        add_text(slide, 0.6, 0.98, 8.0, 0.25, 'SYNTHETIC SMOKE-TEST PREVIEW — figures are NOT from ROT-NOM', size=10,
                 color=RED, bold=True, name='Preview stamp')


def content(prs, L, title, notes_text, fig=None, eqs=(), take=None, cite=None, layout='figL', movie=None):
    s = prs.slides.add_slide(L['content'])
    set_title(s, title)
    strip_empty_placeholders(s)
    if layout == 'figL':
        if movie:
            mw, mh = 5.5, 5.5 * 9 / 16
            add_movie(s, os.path.join(FIG, movie + '.mp4'), os.path.join(FIG, movie + '_poster.png'), FX, FY + 0.15, mw, mh,
                      name='Animation ' + movie)
        elif fig:
            sv, pn = fig_paths(fig)
            add_svg(s, sv, pn, FX, FY, w=FW, h=FH, name='Figure ' + fig)
        y = 1.3
        for i, (k, tex) in enumerate(eqs):
            y += place_eq(s, k, tex, CX, y, CW) + 0.22
        if take:
            add_text(s, CX, y + 0.05, CW, 1.2, take, size=14, color=DARK, name='Takeaway')
    elif layout == 'wide':
        y = 1.2
        for k, tex in eqs:
            y += place_eq(s, k, tex, WX, y, WW) + 0.12
        if fig:
            sv, pn = fig_paths(fig)
            import PIL.Image
            im = PIL.Image.open(pn)
            ar = im.height / im.width
            w = WW
            h = w * ar
            if y + h > 4.45:
                h = 4.45 - y; w = h / ar
            add_svg(s, sv, pn, WX + (WW - w) / 2, y, w=w, h=h, name='Figure ' + fig)
            y += h
        if take:
            add_text(s, WX, min(y + 0.08, 4.25), WW, 0.35, take, size=14, color=DARK, name='Takeaway')
    if cite:
        cy = 4.2 if layout == 'figL' else 4.48
        cx, cw = (CX, CW) if layout == 'figL' else (WX, 6.0)
        add_text(s, cx, cy, cw, 0.35, cite, size=10, color=GREY, name='Citation')
    stamp(s)
    notes(s, notes_text)
    return s


def section(prs, L, title, sub):
    s = prs.slides.add_slide(L['section'])
    for ph in s.placeholders:
        t = ph.placeholder_format.type
        if 'SUBTITLE' in str(t):
            ph.text_frame.text = sub
            for r in ph.text_frame.paragraphs[0].runs:
                r.font.italic = True
        elif 'TITLE' in str(t):
            ph.text_frame.text = title
    strip_empty_placeholders(s)
    return s


def fmt(v, f='.2f'):
    try:
        return format(v, f)
    except Exception:
        return str(v)


def build():
    prs, L = open_template(TEMPLATE)

    # ------------------------------------------------------------ title
    s = prs.slides.add_slide(L['title'])
    for ph in s.placeholders:
        t = str(ph.placeholder_format.type)
        if 'SUBTITLE' in t:
            ph.text_frame.text = 'Event-based structure from motion of a spinning target,\nstage by stage, on the ROT-NOM recording'
            for para in ph.text_frame.paragraphs:
                pPr = para._p.get_or_add_pPr(); pPr.set('marL', '0'); pPr.set('indent', '0')
        elif 'TITLE' in t:
            ph.text_frame.text = 'The ESFM Pipeline'
    strip_empty_placeholders(s)
    add_text(s, 3.1, 4.87, 2.94, 0.28, 'Spacecraft Robotics Laboratory', size=12, color=BLACK, align='c', name='Footer')
    notes(s, 'This deck walks through the ESFM pipeline one idea per slide. Every figure and every number comes from one run '
             'of the pipeline on the ROT-NOM recording from the dataset at '
             'https://www.frdr-dfdr.ca/repo/dataset/1c83d3dc-3a8c-408b-a6c4-9251c0ec1d05, using the standalone Python port of the '
             'MATLAB code (python/ in the repository).')

    # ------------------------------------------------------------ overview
    content(prs, L, 'Eight stages turn events into a point cloud',
            f'The pipeline reads {O.get("n_events", 0)/1e6:.1f} million events covering {fmt(O.get("dur"), ".0f")} s. The front end '
            f'(undistortion and local plane fits) keeps {100*O.get("frac_accept", 0):.0f} % of events and attaches a normal-flow vector to each. '
            f'The warm start tracks edge fragments over the first {6} s and fits a constant spin, which gives {fmt(O.get("w_deg"))} deg/s. '
            f'A first bundle adjustment turns {O.get("n_tracks")} tracks into seed landmarks. Then {O.get("n_slices")} slices of 0.1 s '
            f'are processed causally, growing the map to {O.get("map_size")} landmarks. A final offline refinement keeps '
            f'{O.get("n_final")} landmarks at {fmt(O.get("final_rms"), ".3f")} px normal RMS. The red boxes are the three ideas the rest of the deck '
            'spends the most time on: plane fits, the twist fit, and the progressive loop.',
            fig='F01_overview', layout='wide',
            take='Each red stage introduces a new estimation idea; the grey stages connect them.')

    # ------------------------------------------------------------ 1 events
    section(prs, L, '1 · Events and the camera', 'What the sensor gives us and how pixels become rays')
    content(prs, L, 'An event camera reports changes, not frames',
            'Each pixel fires on its own whenever its log intensity changes by a contrast threshold C since its last event. The output '
            'is a sparse stream of (x, y, t, p) tuples with microsecond timestamps; p is the sign of the change. Nothing happens where '
            'nothing moves, so a slowly spinning target produces events only along its edges and texture. The animation plays 33 ms '
            'of events per frame in real time; red is ON (brighter), blue is OFF (darker). '
            'References: Lichtsteiner, Posch and Delbruck, IEEE JSSC 43(2), 2008; Gallego et al., "Event-based Vision: A Survey", '
            'IEEE TPAMI 44(1), 2022.',
            movie='A1_events',
            eqs=[('ev_model', r'$\displaystyle e_k=(x_k,\,y_k,\,t_k,\,p_k)$'),
                 ('ev_fire', r'$\displaystyle \bigl|\log I(\mathbf{x}_k,t_k)-\log I(\mathbf{x}_k,t_{\mathrm{last}})\bigr|\ \ge\ C$'),
                 ('ev_pol', r'$\displaystyle p_k=\operatorname{sign}\bigl(\Delta\log I\bigr)\in\{-1,+1\}$')],
            take='Static background is silent; only moving edges speak.',
            cite='Lichtsteiner et al., JSSC 2008\nGallego et al., TPAMI 2022')
    f3 = fx.get('f_undistort', {})
    content(prs, L, 'Fisheye distortion is removed per event',
            'The DVXplorer Mini uses a fisheye lens calibrated with the Kannala-Brandt equidistant model. For each event we normalise '
            'the pixel, solve the odd polynomial for the undistorted angle with eight Newton iterations (theta_d is the distorted '
            'angle), and reproject through an ideal pinhole so all later geometry is linear in the normalised coordinates. '
            f'The figure shows how far each pixel moves: k1 = {fmt(f3.get("k1"), ".4f")}, focal = {fmt((f3.get("focal") or [0])[0], ".1f")} px; '
            f'the shift reaches {fmt(f3.get("max_shift_px"), ".0f")} px in the corners because a pinhole stretches wide angles (tan theta). '
            'Reference: Kannala and Brandt, IEEE TPAMI 28(8), 2006.',
            fig='F03_undistortion',
            eqs=[('ud1', r'$\displaystyle \theta_d=\left\|\tfrac{\mathbf{x}-\mathbf{c}}{\mathbf{f}}\right\|=\theta\,(1+k_1\theta^2+k_2\theta^4)$'),
                 ('ud2', r'$\displaystyle \theta\leftarrow\theta-\frac{\theta(1+k_1\theta^2+k_2\theta^4)-\theta_d}{1+3k_1\theta^2+5k_2\theta^4}$'),
                 ('ud3', r'$\displaystyle \mathbf{x}_u=\mathbf{f}\odot\mathbf{m}_d\,\frac{\tan\theta}{\theta_d}+\mathbf{c}$')],
            take='All downstream geometry uses ideal pinhole pixels.',
            cite='Kannala & Brandt, TPAMI 2006')

    # ------------------------------------------------------------ 2 planes
    section(prs, L, '2 · Local planes and normal flow', 'Turning raw events into moving edges')
    content(prs, L, 'In space-time, a moving edge leaves a sheet',
            f'Plotting a short window of events in (x, t, y) shows that an edge moving across the sensor sweeps a surface: at any '
            f'instant it is a line, and as time passes the line slides. The tilt of that surface encodes the edge speed. '
            f'The figure shows 0.3 s starting at t = {fmt(g("f_event_stream", "t_window"))} s (sub-sampled for display); the event rate there '
            f'is about {g("f_event_stream", "rate")/1e3:.0f} k events/s. This picture motivates fitting local planes in the next slide. '
            'Reference: Benosman et al., IEEE TNNLS 25(2), 2014.',
            fig='F02_event_stream',
            eqs=[('sheet', r'$\displaystyle \mathbf{x}(t)\approx\mathbf{x}_0+\mathbf{v}\,(t-t_0)\ \ \text{along the edge normal}$')],
            take='Edges become planes; plane tilt is speed.',
            cite='Benosman et al., TNNLS 2014')
    f4 = fx.get('f_plane_fit', {})
    content(prs, L, 'A plane fit is an eigen-decomposition',
            'Events are grouped into cells of 8 x 8 px x 0.8 s, separately per polarity. Time is scaled by sigma_t = 0.0375 cell_s '
            '(0.03 s) so a pixel and sigma_t seconds have equal length; this keeps the shape thresholds independent of the cell '
            'duration. The 3 x 3 covariance of the scaled events is decomposed: the smallest eigenvector is the plane normal, the '
            'smallest eigenvalue is the sheet thickness and the middle one its extent along the edge. '
            f'The cell shown has {f4.get("cell_events")} events with eigenvalues {", ".join(fmt(v) for v in f4.get("eig", []))}. '
            'Reference: Benosman et al., IEEE TNNLS 2014 (local plane fitting for event-based flow).',
            fig='F04_plane_fit',
            eqs=[('pf1', r'$\displaystyle \mathbf{p}_i=\bigl(x_i,\;y_i,\;t_i/\sigma_t\bigr)^{\!\top}$'),
                 ('pf2', r'$\displaystyle \boldsymbol{\Sigma}=\frac1n\sum_i(\mathbf{p}_i-\bar{\mathbf{p}})(\mathbf{p}_i-\bar{\mathbf{p}})^{\!\top}$'),
                 ('pf3', r'$\displaystyle \boldsymbol{\Sigma}\mathbf{v}_j=\lambda_j\mathbf{v}_j,\quad \mathbf{n}=\mathbf{v}_1\ (\lambda_1\le\lambda_2\le\lambda_3)$')],
            take='Smallest eigenvector = plane normal.',
            cite='Benosman et al., TNNLS 2014')
    f5 = fx.get('f_normal_flow', {})
    content(prs, L, 'The plane tilt gives the normal flow',
            'The plane x n_x + y n_y + (t/sigma_t) n_t = const has time gradient g = -sigma_t (n_x, n_y)/n_t in seconds per pixel. '
            'Inverting the slowness gives the velocity v = g/|g|^2 in pixels per second. Only the component across the edge is '
            'recoverable (the aperture problem), which is why every later residual is measured along the edge normal. '
            f'Arrows are averaged per 8 px block and drawn as one second of motion; the median speed in this slice is '
            f'{fmt(f5.get("median_speed"), ".1f")} px/s.',
            fig='F05_normal_flow',
            eqs=[('nf1', r'$\displaystyle \mathbf{g}=-\,\sigma_t\,\frac{(n_x,\,n_y)}{n_t}\quad[\mathrm{s/px}]$'),
                 ('nf2', r'$\displaystyle \mathbf{v}=\frac{\mathbf{g}}{\|\mathbf{g}\|^{2}}\quad[\mathrm{px/s}]$')],
            take='Only motion across the edge is observable.',
            cite='Aperture problem; Benosman et al., TNNLS 2014')
    f6 = fx.get('f_accept', {})
    content(prs, L, 'Three tests keep only clean, moving sheets',
            'A cell is accepted if (1) it is thin: lambda1/lambda2 <= 0.45; (2) it moves: |n_t| > 0.05, which rejects static or '
            'flickering pixels whose normal lies in the image plane; and (3) warping its events back to the cell time along the '
            'fitted flow piles at least three of them on the same pixel. Two grids offset by half a cell in x, y and t are fitted and '
            'each event keeps the more planar of its two cells, so edges on a cell boundary are not lost. '
            f'Over the whole run {100*f6.get("frac_events", 0):.1f} % of events pass. The scatter shows the {f6.get("cells")} ON-polarity cells of 3.2 s.',
            fig='F06_accept',
            eqs=[('acc', r'$\displaystyle \frac{\lambda_1}{\lambda_2}\le 0.45,\qquad |n_t|>0.05,\qquad \text{support}\ge 3$')],
            take='Rejected: blobs, flicker, and cells too sparse to fit.')
    f7 = fx.get('f_warp', {})
    content(prs, L, 'Warping by the flow sharpens every slice',
            'Within a 0.1 s slice each event is moved along its own normal flow to the slice centre t_c. Smear across the edge '
            'collapses, so edges become thin lines at a single instant; along the edge nothing changes because normal flow has no '
            'tangential component. This is the same principle as contrast maximisation (Gallego, Rebecq and Scaramuzza, CVPR 2018), '
            'but here the motion model is per event (local normal flow) instead of one global warp. '
            f'Image variance in the zoomed window goes from {fmt(f7.get("var_raw"))} to {fmt(f7.get("var_warp"))}.',
            fig='F07_warp',
            eqs=[('warp', r'$\displaystyle \mathbf{x}_k^{\prime}=\mathbf{x}_k-\mathbf{v}_k\,(t_k-t_c)$')],
            take='Sharper edges make association unambiguous.',
            cite='Gallego et al., CVPR 2018')

    # ------------------------------------------------------------ 3 warm start
    section(prs, L, '3 · Warm start', 'A constant twist explains the first six seconds')
    content(prs, L, 'An alpha-beta filter tracks edge fragments',
            'In each slice, every live track predicts its position with constant velocity, claims the unclaimed warped events within '
            '3 px, and updates with fixed gains. Beta follows the Benedict-Bordner relation beta = alpha^2/(2 - alpha) with alpha = 0.6. '
            'Unclaimed events are binned into 6 px cells; a cell with at least 3 events seeds a new track. Each sample also stores the '
            f'mean normal-flow direction, which becomes the edge normal of the observation. {O.get("n_tracks")} tracks of at least 20 samples survive. '
            'Related work to check the novelty claim in track_events.m against: Gehrig et al., "EKLT", IJCV 128, 2020. '
            'Reference: Benedict and Bordner, IRE Trans. Automatic Control AC-7(4), 1962.',
            movie='A2_tracks',
            eqs=[('ab1', r'$\displaystyle \hat{\mathbf{p}}=\mathbf{p}+\mathbf{v}\,\Delta t$'),
                 ('ab2', r'$\displaystyle \mathbf{p}\leftarrow\hat{\mathbf{p}}+\alpha(\mathbf{z}-\hat{\mathbf{p}}),\quad \mathbf{v}\leftarrow\mathbf{v}+\tfrac{\beta}{\Delta t}(\mathbf{z}-\hat{\mathbf{p}})$'),
                 ('ab3', r'$\displaystyle \beta=\frac{\alpha^2}{2-\alpha},\qquad \alpha=0.6$')],
            take='Tracks carry position and edge normal.',
            cite='Benedict & Bordner, IRE TAC 1962')
    f9 = fx.get('f_obs_matrix', {})
    content(prs, L, 'Each track sample becomes one observation',
            'Tracks start and stop at different slices, so they are aligned on the common slice grid. Each filled cell is one '
            'observation: the slice time, the track id, the pixel position and the unit edge normal. Missing samples are simply absent; '
            f'no interpolation is done. In the warm window {O.get("n_warm_obs")} observations fill {100*f9.get("fill", 0):.0f} % of the slots.',
            fig='F09_obs_matrix',
            eqs=[('obs', r'$\displaystyle \mathbf{o}_{ij}=\bigl(t_j,\ i,\ u_{ij},\ v_{ij},\ n^x_{ij},\ n^y_{ij}\bigr)$')],
            take='Observation = where an edge was, and which way it faced.')
    content(prs, L, 'A constant twist: one axis, one rate',
            'For the warm window the target is modelled as spinning at a constant rate omega about a fixed unit axis a (both in the '
            'camera frame at t0), with its origin at T = [c_x, c_y, 1]. Fixing the depth of the origin at 1 sets the scale gauge, so '
            'the reconstruction is in units of the initial range. Rodrigues gives R(t); the pose maps object points into the camera. '
            f'The fitted values: |omega| = {fmt(g("f_twist", "w_deg"))} deg/s, a = [{", ".join(fmt(v, ".3f") for v in g("f_twist", "a", []))}].',
            fig='F10_twist',
            eqs=[('tw1', r'$\displaystyle \mathbf{R}(t)=\mathbf{I}+\sin\theta\,[\mathbf{a}]_\times+(1-\cos\theta)[\mathbf{a}]_\times^2,\ \ \theta=\omega(t-t_0)$'),
                 ('tw2', r'$\displaystyle \mathbf{T}=[c_x,\ c_y,\ 1]^{\top},\qquad \mathbf{Y}=\mathbf{R}\mathbf{X}+\mathbf{T}$')],
            take='Five unknowns: axis (2), rate, centre (2).')
    content(prs, L, 'Triangulation weights the edge normal',
            'Given a candidate motion, each track is triangulated linearly. A pinhole observation gives two DLT rows. They are '
            'recombined along the edge normal n (weight 1) and the tangent t (weight epsilon = 0.15), because an edge pins the point '
            'across the edge but only weakly along it. Accumulating the 3 x 3 normal equations per track and solving gives the point. '
            f'The rays shown come from track {g("f_rays", "track")}: {g("f_rays", "n_rays")} observations spanning {fmt(g("f_rays", "turn_deg"), ".0f")} deg of rotation, drawn in the object frame. '
            'Reference: Hartley and Zisserman, Multiple View Geometry, 2nd ed., 2004.',
            fig='F11_rays',
            eqs=[('tr1', r'$\displaystyle \mathbf{a}_x=m_x\mathbf{r}_3-\mathbf{r}_1,\ \ b_x=T_x-m_xT_z\ \ (\text{same for }y)$'),
                 ('tr2', r'$\displaystyle \bigl[\mathbf{n}^{\!\top};\ \varepsilon\,\mathbf{t}^{\!\top}\bigr]\,[\mathbf{a}_x\ \mathbf{a}_y]^{\!\top}\mathbf{X}=\bigl[\mathbf{n}^{\!\top};\ \varepsilon\,\mathbf{t}^{\!\top}\bigr][b_x\ b_y]^{\!\top}$'),
                 ('tr3', r'$\displaystyle \mathbf{X}_i=\Bigl(\sum_j\mathbf{A}_{ij}^{\!\top}\mathbf{A}_{ij}\Bigr)^{-1}\sum_j\mathbf{A}_{ij}^{\!\top}\mathbf{b}_{ij}$')],
            take='Across-edge rows dominate; along-edge rows only regularise.',
            cite='Hartley & Zisserman, 2004')
    content(prs, L, 'The error is measured across the edge',
            'The quantity every fit minimises is the signed distance between the observed and projected position, measured along '
            'the observed edge normal. Moving the projection along the edge costs nothing, which matches what the events can actually '
            'tell us. In the warm fit, an observation whose point falls closer than 0.2 depth units gets a fixed penalty of 10 px. '
            f'This track ({g("f_normal_residual", "track")}) has an RMS normal residual of {fmt(g("f_normal_residual", "rms"))} px at the warm-start motion.',
            fig='F12_normal_residual',
            eqs=[('nr1', r'$\displaystyle r_{ij}=\mathbf{n}_{ij}^{\!\top}\bigl(\mathbf{u}_{ij}-\pi(\mathbf{R}_j\mathbf{X}_i+\mathbf{T}_j)\bigr)$'),
                 ('nr2', r'$\displaystyle \pi(\mathbf{Y})=\mathbf{f}\odot\frac{(Y_x,\,Y_y)}{Y_z}+\mathbf{c}$')],
            take='Point-to-line distance, not point-to-point.')
    content(prs, L, 'A Cauchy loss caps the influence of outliers',
            'Residuals pass through the Cauchy loss with scale c = 4 px. For small r it is quadratic; for large r it grows only '
            'logarithmically, so a wrong association cannot dominate. In iteratively reweighted form it is a per-residual weight '
            f'1/(1 + r^2/c^2). {100*g("f_cauchy", "frac_out", 0):.1f} % of the seed residuals lie beyond 2c and get weight below 0.2. '
            'Reference: Barron, "A General and Adaptive Robust Loss Function", CVPR 2019 (Cauchy is its alpha = 0 case); '
            'Triggs et al., "Bundle Adjustment - A Modern Synthesis", 2000.',
            fig='F13_cauchy',
            eqs=[('ca1', r'$\displaystyle \rho(r)=c^2\log\!\Bigl(1+\frac{r^2}{c^2}\Bigr)$'),
                 ('ca2', r'$\displaystyle w(r)=\frac{\rho^{\prime}(r)}{2r}=\frac{1}{1+r^2/c^2},\qquad c=4\ \mathrm{px}$')],
            take='Quadratic near zero, logarithmic in the tails.',
            cite='Barron, CVPR 2019\nTriggs et al., 2000')
    content(prs, L, 'A sweep over 26 axes finds the spin',
            'The cost is non-convex in the axis, so the axis is not optimised from one start. Thirteen directions on a Fibonacci '
            'hemisphere and their negatives give 26 fixed axes; for each, and for five starting rates (2 to 10 deg/s), the rate and '
            'centre are fitted with bounded nonlinear least squares on the Cauchy-transformed residuals (structure is re-triangulated '
            f'inside every evaluation). That is {g("f_sweep", "n_fits")} fits. The best one then frees the axis too. '
            'Reference: Gonzalez, Mathematical Geosciences 42(1), 2010 (Fibonacci lattice).',
            fig='F14_sweep',
            eqs=[('sw1', r'$\displaystyle z_i=1-\tfrac{i+1/2}{n},\quad \phi_i=\pi(1+\sqrt5)\,i,\quad n=13$'),
                 ('sw2', r'$\displaystyle \min_{\omega,\,\mathbf{c}}\ \sum_{ij}\rho\bigl(r_{ij}(\mathbf{a}_k,\omega,\mathbf{c})\bigr)\quad \forall\,\mathbf{a}_k$')],
            take='Coarse global search, then local refinement.',
            cite='González, Math. Geosci. 2010')
    f15 = fx.get('f_rate_profile', {})
    content(prs, L, 'The rate profile confirms a single minimum',
            'With the refined axis fixed, the rate is stepped from 1 to 12 deg/s and the centre re-fitted at each value. The profile '
            'shows how well the rate is determined and whether another rate could explain the tracks nearly as well. '
            f'The refined rate is {fmt(f15.get("w"))} deg/s; costs within 2 of the minimum span {fmt((f15.get("band") or [0, 0])[0], ".1f")}-{fmt((f15.get("band") or [0, 0])[1], ".1f")} deg/s.',
            fig='F15_rate_profile',
            eqs=[('rp', r'$\displaystyle C(\omega)=\min_{\mathbf{c}}\ \sum_{ij}\rho\bigl(r_{ij}(\mathbf{a}^{\star},\omega,\mathbf{c})\bigr)$')],
            take='One clear valley: the rate is identifiable from six seconds.')

    # ------------------------------------------------------------ 4 trajectory + BA
    section(prs, L, '4 · Continuous-time trajectory', 'A spline on top of the twist, solved by bundle adjustment')
    content(prs, L, 'Spin = constant rate + smooth spline',
            'Over two minutes the spin is not exactly constant, so the angle and the translation get a uniform cubic B-spline '
            f'correction with 1 s knot spacing ({O.get("K")} coefficients each). The constant rate keeps the bulk of the motion; the spline '
            'absorbs slow variations. Every time instant is influenced by exactly four coefficients, which keeps the Jacobian sparse. '
            'References: Furgale, Barfoot and Sibley, ICRA 2012; Lovegrove, Patron-Perez and Sibley, BMVC 2013; '
            'Mueggler et al., IEEE T-RO 34(6), 2018 (continuous-time trajectories for event cameras).',
            fig='F16_spline',
            eqs=[('sp1', r'$\displaystyle \theta(t)=\omega\,(t-t_0)+\sum_j k_j\,B_j(t)$'),
                 ('sp2', r'$\displaystyle \mathbf{T}(t)=[c_x,c_y,1]^{\top}+\sum_j \mathbf{m}_j\,B_j(t)$'),
                 ('sp3', r'$\displaystyle B(s)=\tfrac16\bigl[(1-s)^3,\,3s^3-6s^2+4,\,\dots,\,s^3\bigr]$')],
            take='Four active coefficients at any instant.',
            cite='Furgale et al., ICRA 2012\nMueggler et al., T-RO 2018')
    content(prs, L, 'One sparse, robust least-squares problem',
            'Bundle adjustment stacks: normal and tangent image rows (tangent weighted 0.15) under the Cauchy loss; squared priors on '
            'spline curvature (second differences divided by h^2) and a weak pull toward zero correction; gauge rows forcing the '
            'correction to vanish at t0 (so R(t0) = I and the depth is 1); a row pinning the landmark centroid along the axis (a '
            'shift along the axis is unobservable); and a hinge on depth below 0.2. The sparsity plot shows the structure for '
            f'60 seed landmarks: dense motion columns, block-diagonal landmark columns. '
            'References: Triggs et al. 2000; depth hinge after Parameshwara et al., DiffPoseNet, CVPR 2022.',
            fig='F17_jacobian',
            eqs=[('ba1', r'$\displaystyle E=\sum\rho(r^{n})+\sum\rho(\varepsilon r^{t})+\|\mathbf{r}_{\mathrm{prior}}\|^2+\|\mathbf{r}_{\mathrm{depth}}\|^2$'),
                 ('ba2', r'$\displaystyle r_{\mathrm{curv}}=\lambda\,\frac{k_j-2k_{j+1}+k_{j+2}}{h^2},\quad \sum_j B_j(t_0)\,k_j=0$'),
                 ('ba3', r'$\displaystyle \mathbf{a}^{\!\top}\bar{\mathbf{X}}=0,\qquad r_{\mathrm{depth}}=w\max(0,\,0.2-Y_z)$')],
            take='Priors fix the gauge; the loss handles outliers.',
            cite='Triggs et al., 2000')
    content(prs, L, 'Levenberg-Marquardt with robust weights',
            'Each iteration linearises analytically (verified against finite differences), applies the Cauchy weights as row scales '
            'so the weighted gradient equals the robust gradient, and solves the damped normal equations with Marquardt scaling. '
            'A step is accepted only if the true robust cost decreases; damping drops by 3 on success and rises by 4 on failure. '
            f'For the seed map the normal RMS goes from {fmt(O.get("seed_rms0"))} px (linear triangulation) to {fmt(O.get("seed_rms1"))} px '
            '(the solver minimises the robust cost including tangent rows and priors, not the normal RMS alone). '
            'References: Levenberg 1944; Marquardt 1963.',
            fig='F18_lm',
            eqs=[('lm1', r'$\displaystyle (\mathbf{J}^{\!\top}\mathbf{W}\mathbf{J}+\lambda\mathbf{D})\,\boldsymbol{\delta}=-\mathbf{J}^{\!\top}\mathbf{W}\mathbf{r}$'),
                 ('lm2', r'$\displaystyle \mathbf{W}=\mathrm{diag}\,w(r),\quad \mathbf{D}=\mathrm{diag}(\mathbf{J}^{\!\top}\mathbf{W}\mathbf{J})$')],
            take='Accept only steps that lower the true cost.',
            cite='Levenberg 1944; Marquardt 1963')

    # ------------------------------------------------------------ 5 progressive
    section(prs, L, '5 · Progressive mapping', 'One 0.1 s slice at a time, using only past data')
    f19 = fx.get('f_cycle', {})
    content(prs, L, 'Every slice runs the same seven steps',
            'After the warm start the recording is replayed as a stream of 0.1 s slices. Each slice warps its events, predicts where '
            'every landmark should be, lets landmarks claim events, corrects the recent motion, re-associates with the corrected '
            'motion, commits the claims, and then grows 2-D candidate tracks on the leftover events and promotes mature ones to '
            f'landmarks. The numbers are from the slice at t_c = {fmt(f19.get("t"))} s.',
            fig='F19_cycle', layout='wide',
            take='Causal: each slice sees only the map built from earlier data.')
    content(prs, L, 'Prediction: where each landmark should be',
            'The current trajectory is evaluated at the slice centre; every landmark is projected, and its image velocity is the '
            'central difference of the projection over +/- 0.05 s. A landmark is a candidate for association if it is in front of '
            f'the camera (depth > 0.2) and inside the image plus an 8 px margin. In this slice {g("f_predict", "n_in")} landmarks are in view.',
            fig='F20_predict',
            eqs=[('pr1', r'$\displaystyle \hat{\mathbf{u}}_i=\pi\bigl(\mathbf{R}(t_c)\mathbf{X}_i+\mathbf{T}(t_c)\bigr)$'),
                 ('pr2', r'$\displaystyle \dot{\hat{\mathbf{u}}}_i\approx\frac{\hat{\mathbf{u}}_i(t_c+h)-\hat{\mathbf{u}}_i(t_c-h)}{2h},\quad h=0.05\,\mathrm{s}$')],
            take='The map tells the events where to look.')
    content(prs, L, 'Four gates decide which events a landmark claims',
            'Around each prediction, unclaimed warped events within 8 px are tested in the frame of each event\'s own flow direction '
            'd: within 3 px across the edge, within 8 px along it, moving the predicted way, and with an orientation within 60 deg of '
            'the landmark\'s last normal. Landmarks with more observations claim first, and an event can be claimed once. A claim '
            'becomes one observation built from the raw (unwarped) event positions at their mean time. '
            f'Here {g("f_gates", "ok")} of {g("f_gates", "cand")} nearby events pass.',
            fig='F21_gates',
            eqs=[('g1', r'$\displaystyle |\mathbf{d}_e^{\!\top}\boldsymbol{\Delta}_e|<3,\qquad |\mathbf{d}_e\times\boldsymbol{\Delta}_e|<8$'),
                 ('g2', r'$\displaystyle \mathbf{d}_e^{\!\top}\dot{\hat{\mathbf{u}}}_i>0,\qquad |\mathbf{d}_e^{\!\top}\mathbf{n}_i|>0.5$')],
            take='Each claim becomes one new observation.')
    content(prs, L, 'Only the recent knots move during correction',
            'After the first association, a motion-only adjustment frees just the spline coefficients that influence the last '
            '3 s (three knot spacings) and keeps the mean motion and every landmark fixed; it uses the last 3 s of observations plus '
            'this slice\'s claims and runs 3 LM iterations. The newest coefficient is constrained only by the curvature prior, '
            'so it extrapolates smoothly into the next slice. The right panel shows how far predictions move between the two passes.',
            fig='F22_correction',
            eqs=[('cm', r'$\displaystyle \mathcal{F}=\bigl\{k_j,\mathbf{m}_j:\ B_j\neq0\ \text{on}\ [t_c-3h,\ t_c]\bigr\}$')],
            take='Bounded cost per slice: a handful of parameters.')
    content(prs, L, 'Candidates become landmarks once they turn',
            'Events no landmark claimed feed 2-D candidate tracks (the same alpha-beta tracker). A candidate is mature when it has '
            'at least 12 samples and the target has turned at least 8 deg between its first and last sample, which guarantees a '
            'viewing baseline. It is triangulated with the current trajectory and promoted if it lies in front of the camera and '
            f'reprojects within 2 px RMS. Over the run {g("f_growth", "promoted")} landmarks were promoted.',
            fig='F23_growth',
            eqs=[('pm1', r'$\displaystyle S\ge12\quad\text{and}\quad |\theta(t_S)-\theta(t_1)|\ge 8^{\circ}$'),
                 ('pm2', r'$\displaystyle \min_j Y_{z,j}>0.2\quad\text{and}\quad \mathrm{RMS}_j(r)<2\ \mathrm{px}$')],
            take='Baseline first, then triangulate.')
    content(prs, L, 'The map grows as the target turns',
            'Snapshots of the map once per second during the progressive loop, in model units and viewed with the spin axis vertical. '
            'Red points were promoted in the preceding second. New structure appears as previously unseen faces rotate into view.',
            movie='A3_growth',
            eqs=[('gr', r'$\displaystyle N(t)=N_{\mathrm{seed}}+\sum_{t_c\le t}n_{\mathrm{promoted}}(t_c)$')],
            take='Coverage follows the rotation.')
    content(prs, L, 'Local and global adjustments keep the map consistent',
            'A local adjustment runs at each new knot (every 1 s): the spline coefficients of the last 3 s and every landmark seen '
            'in that window move, using all observations of those landmarks so their older history anchors them. A global '
            'adjustment (mean motion, all released knots, all landmarks) runs whenever the observation count has grown by 25 %, so '
            f'it is frequent early and rare later. This run did {g("f_schedule", "n_local")} local and {g("f_schedule", "n_global")} global adjustments. '
            'The local/global split follows keyframe SLAM practice: Mouragnon et al., CVPR 2006; Klein and Murray, ISMAR 2007.',
            fig='F24_schedule',
            eqs=[('sc', r'$\displaystyle \text{global when}\quad N_{\mathrm{obs}}\ \ge\ 1.25\,N_{\mathrm{obs}}^{\mathrm{last}}$')],
            take='Local every second; global when the map has grown.',
            cite='Mouragnon et al., CVPR 2006\nKlein & Murray, ISMAR 2007')
    content(prs, L, 'The online estimate tracks the final one',
            'The upper panel compares the rate the causal loop reported at each slice with the final offline trajectory; the lower '
            'panel shows the angle difference. The two share a gauge only approximately (the global adjustments can shift it), so '
            f'small offsets are expected. The online angle is within {fmt(g("f_online", "err95"))} deg of the final angle 95 % of the time.',
            fig='F25_online',
            eqs=[('on', r'$\displaystyle \Delta\theta(t_c)=\theta_{\mathrm{online}}(t_c)-\theta_{\mathrm{final}}(t_c)$')],
            take='Real-time motion is close to the smoothed answer.')
    content(prs, L, 'The map stays locked onto the events',
            'Twelve seconds of accepted events with the final landmarks projected through the refined trajectory. Agreement between '
            'the red circles and the grey edges is a direct visual check of motion and structure together.',
            movie='A5_reprojection',
            eqs=[('lk', r'$\displaystyle \hat{\mathbf{u}}_i(t)=\pi\bigl(\mathbf{R}(t)\mathbf{X}_i+\mathbf{T}(t)\bigr)$')],
            take='Motion and structure agree across the run.')

    # ------------------------------------------------------------ 6 result
    section(prs, L, '6 · Result', 'Final refinement and what it produced')
    f26 = fx.get('f_filter', {})
    content(prs, L, 'Final refinement drops poorly fitting points',
            'Offline, everything is freed over the whole run: mean motion, all spline coefficients and every landmark with at '
            'least 20 observations. After one pass, landmarks whose normal RMS across their observations is 2 px or more are dropped '
            f'and the survivors are refined again. Of {f26.get("eligible")} eligible landmarks, {f26.get("kept")} are kept and {f26.get("dropped")} dropped.',
            fig='F26_filter',
            eqs=[('fr', r'$\displaystyle \mathrm{RMS}_i=\sqrt{\tfrac{1}{|\mathcal{O}_i|}\sum_{j\in\mathcal{O}_i}r_{ij}^2}\ <\ 2\ \mathrm{px}$')],
            take='Two passes: refine, filter, refine.')
    content(prs, L, 'The reconstructed point cloud',
            f'{O.get("n_final")} landmarks in metres, centred and rotated so the spin axis is vertical; colour is the number of observations. '
            'The scale is provisional: without known endpoints, the 1-99 percentile extent along the cloud\'s principal axis is '
            f'assumed to equal the known 1.27 m panel width, giving {fmt(O.get("scale"), ".3f")} m per model unit. Scaling never changes '
            'the shape or any reprojection.',
            fig='F27_cloud', layout='wide',
            eqs=[('scale', r'$\displaystyle s=\frac{W}{p_{99}-p_{1}},\qquad W=1.27\ \mathrm{m}\quad(\text{provisional})$')])
    content(prs, L, 'The cloud from every side',
            'Turntable view of the final cloud about the spin axis.',
            movie='A4_cloud',
            eqs=[('tt', r'$\displaystyle \mathbf{X}_i^{\mathrm{m}}=s\,\mathbf{X}_i$')],
            take=f'{O.get("n_final")} landmarks, {O.get("final_obs", 0):,} observations.')
    f28 = fx.get('f_residuals', {})
    content(prs, L, 'Residuals stay near one pixel for the whole run',
            'Distribution of the final normal residuals per second of the recording with the per-second RMS overlaid. A flat band '
            'means the trajectory model holds over the full duration, with no drift in the spline. '
            f'Overall RMS {fmt(f28.get("rms"), ".3f")} px over {f28.get("n", 0):,} observations.',
            fig='F28_residuals',
            eqs=[('rms', r'$\displaystyle \mathrm{RMS}=\sqrt{\tfrac{1}{M}\sum_{ij}r_{ij}^2}$')],
            take=f'{fmt(f28.get("rms"), ".2f")} px RMS across the edge.')
    content(prs, L, 'Most events are explained by the map',
            'Fraction of each slice\'s accepted events claimed by existing landmarks. The remainder feeds the candidate tracker; a '
            'high fraction means the map accounts for the scene the camera sees.',
            fig='F29_explained',
            eqs=[('ex', r'$\displaystyle \eta(t_c)=\frac{|\{e\ \text{claimed}\}|}{|\{e\ \text{accepted in slice}\}|}$')],
            take=f'Median {fmt(g("f_explained", "median"), ".0f")} % per slice.')

    # ------------------------------------------------------------ summary + end
    content(prs, L, 'Summary',
            f'Recap with the run\'s numbers: {O.get("n_events", 0)/1e6:.1f} M events, {100*O.get("frac_accept", 0):.0f} % accepted, '
            f'{O.get("n_tracks")} warm tracks, a {fmt(O.get("w_deg"))} deg/s twist, {O.get("map_size")} landmarks built causally, '
            f'{O.get("n_final")} kept at {fmt(O.get("final_rms"), ".3f")} px. Final mean spin {fmt(O.get("final_w_deg"), ".3f")} deg/s.',
            fig='F01_overview', layout='wide',
            take=f'{O.get("n_final")} landmarks at {fmt(O.get("final_rms"), ".2f")} px from events alone, with a causal loop and one offline refinement.')
    s = prs.slides.add_slide(L['end'])
    s.shapes.title.text_frame.text = 'Thank You!'
    strip_empty_placeholders(s)
    prs.save(OUTFILE)
    print('saved', OUTFILE, len(prs.slides), 'slides')


if __name__ == '__main__':
    build()
